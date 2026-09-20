"""Extrai o que importa de cada planilha da clínica, sem depender de posição fixa."""
import openpyxl, warnings, unicodedata
from collections import defaultdict
warnings.filterwarnings("ignore")

def norm(t):
    t = unicodedata.normalize("NFD", str(t or ""))
    return "".join(c for c in t if unicodedata.category(c) != "Mn").lower().strip()

def achar_cabecalho(ws, chaves, limite=15):
    """Acha a linha do cabeçalho procurando as palavras-chave — em vez de
       fixar 'linha 5', que quebra quando alguém insere uma linha acima."""
    for i, linha in enumerate(ws.iter_rows(max_row=limite, values_only=True), 1):
        vals = [norm(c) for c in linha if c]
        if sum(1 for k in chaves if any(k in v for v in vals)) >= len(chaves) - 1:
            return i, {norm(c): j for j, c in enumerate(linha) if c}
    return None, {}

def ler_lancamentos(arq):
    """A fonte de verdade: cada movimento de dinheiro, com data, valor e conta."""
    wb = openpyxl.load_workbook(arq, data_only=True, read_only=True)
    saida = []
    for nome in wb.sheetnames:
        ws = wb[nome]
        cab_i, cols = achar_cabecalho(ws, ["status", "valor", "contas", "fornecedor"])
        if not cab_i:
            continue
        c = lambda *nomes: next((cols[n] for n in nomes if n in cols), None)
        i_data, i_val = c("data doc", "data"), c("valor")
        i_grupo, i_conta = c("contas"), c("texto")
        i_forn, i_st = c("fornecedor"), c("status")
        i_mes, i_ano = c("mes"), c("ano")
        if i_val is None:
            continue
        for l in ws.iter_rows(min_row=cab_i + 1, values_only=True):
            if len(l) <= i_val or not isinstance(l[i_val], (int, float)):
                continue
            d = l[i_data] if i_data is not None and i_data < len(l) else None
            mes = l[i_mes] if i_mes is not None and i_mes < len(l) else None
            ano = l[i_ano] if i_ano is not None and i_ano < len(l) else None
            if (not mes or not ano) and hasattr(d, "month"):
                mes, ano = d.month, d.year
            if not mes or not ano:
                continue
            saida.append({
                "competencia": f"{int(ano)}-{int(mes):02d}",
                "data": d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else None,
                "valor": float(l[i_val]),
                "grupo": str(l[i_grupo]).strip() if i_grupo is not None and i_grupo < len(l) and l[i_grupo] else "sem grupo",
                "conta": str(l[i_conta]).strip() if i_conta is not None and i_conta < len(l) and l[i_conta] else "",
                "fornecedor": str(l[i_forn]).strip() if i_forn is not None and i_forn < len(l) and l[i_forn] else "",
                "status": str(l[i_st]).strip() if i_st is not None and i_st < len(l) and l[i_st] else "",
                "origem": f"{arq}·{nome}",
            })
    wb.close()
    return saida

# Receita se identifica pelo grupo, não por sinal — na planilha tudo é positivo.
PALAVRAS_RECEITA = ["particular", "convenio", "venda", "receita", "hof"]

def eh_receita(lanc):
    g = norm(lanc["grupo"])
    if "deducao" in g or "deducoes" in g:
        return False
    return any(p in g for p in PALAVRAS_RECEITA)

def consolidar(lancs):
    meses = defaultdict(lambda: {"receita": 0.0, "custo": 0.0, "n": 0})
    grupos = defaultdict(float)
    fornec = defaultdict(float)
    for l in lancs:
        m = meses[l["competencia"]]
        if eh_receita(l):
            m["receita"] += l["valor"]
        else:
            m["custo"] += l["valor"]
            if l["fornecedor"]:
                fornec[l["fornecedor"]] += l["valor"]
        m["n"] += 1
        grupos[l["grupo"]] += l["valor"]
    for m in meses.values():
        m["resultado"] = m["receita"] - m["custo"]
        m["margem"] = (m["resultado"] / m["receita"] * 100) if m["receita"] else 0
    return {"meses": dict(sorted(meses.items())),
            "grupos": dict(sorted(grupos.items(), key=lambda x: -x[1])),
            "fornecedores": dict(sorted(fornec.items(), key=lambda x: -x[1])[:15]),
            "total_lancamentos": len(lancs)}
